from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import yaml


ROOT = Path(__file__).resolve().parents[1]
LESSON = ROOT / "教案生成器" / "lesson-plan-docx-generator"
V112 = LESSON / "assets" / "templates" / "lesson-plan" / "v1.1.2"
V113 = LESSON / "assets" / "templates" / "lesson-plan" / "v1.1.3"
ISSUE71_INPUT = LESSON / "docs" / "evidence" / "issue-71-synthetic-regression" / "synthetic-lesson-content.json"
EXPECTED_V112_TEMPLATE_SHA256 = "6FFAFA579D3AACBC535BA624D0A6A766644868B0AF1FCE6AC37B20BA8F3D8FC1"
EXPECTED_V112_MANIFEST_SHA256 = "31A9A13783E6233235CC80A84E0D111A82F667F1E49A1019C9AFE8EEFC8E5AF9"
OPTION_B_WIDTHS_TWIPS = [1304, 2948, 737, 2769]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def _run(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env.update({"PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"})
    return subprocess.run(
        [sys.executable, str(script), *args],
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def _generate(source: Path, output: Path, template: Path | None = None, legacy: bool = False) -> dict:
    args = ["--tasks-json", str(source), "--output-dir", str(output)]
    if template is not None:
        # Template-only resolution must discover the matching v1.1.3 manifest.
        args.extend(("--template", str(template)))
    if legacy:
        args.append("--legacy")
    result = _run(LESSON / "scripts" / "generate_lesson_plans.py", *args)
    if result.returncode != 0:
        raise AssertionError(result.stderr or result.stdout)
    return json.loads((output / "qa-report.json").read_text(encoding="utf-8"))


def _evaluation_table(path: Path):
    document = Document(path)
    return document, document.tables[0].cell(12, 1).tables[0]


def _evaluation_matrix(path: Path) -> list[list[str]]:
    _document, table = _evaluation_table(path)
    return [[cell.text for cell in row.cells] for row in table.rows]


def _bookmark_signature(path: Path) -> list[tuple[str, str]]:
    with zipfile.ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    return [
        (node.get(qn("w:name"), ""), node.get(qn("w:id"), ""))
        for node in root.xpath(".//w:bookmarkStart", namespaces={"w": W_NS})
    ]


def _merge_signature(path: Path) -> list[list[list[tuple[str, str]]]]:
    _document, table = _evaluation_table(path)
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            properties = cell._tc.tcPr
            cells.append(
                [
                    (child.tag, child.get(qn("w:val"), ""))
                    for child in properties
                    if child.tag in {qn("w:vMerge"), qn("w:gridSpan")}
                ]
                if properties is not None
                else []
            )
        rows.append(cells)
    return rows


class Issue71TemplateV113Tests(unittest.TestCase):
    def test_protected_fixture_and_template_manifests_are_checked_out_as_lf(self) -> None:
        paths = [
            ISSUE71_INPUT.relative_to(ROOT).as_posix(),
            (V112 / "manifest.yaml").relative_to(ROOT).as_posix(),
            *[path.relative_to(ROOT).as_posix() for path in sorted(V113.rglob("manifest.yaml"))],
        ]
        result = subprocess.run(
            ["git", "check-attr", "-z", "eol", "--", *paths],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        fields = result.stdout.split("\0")
        self.assertEqual(fields[-1], "")
        triples = [fields[index : index + 3] for index in range(0, len(fields) - 1, 3)]
        attributes = {path: value for path, name, value in triples if name == "eol"}
        self.assertEqual(attributes, {path: "lf" for path in paths})
        for relative_path in paths:
            self.assertNotIn(b"\r\n", (ROOT / relative_path).read_bytes(), relative_path)

    def test_v113_has_declared_evaluation_layout_and_v112_is_unchanged(self) -> None:
        old_template = V112 / "template.docx"
        old_manifest = V112 / "manifest.yaml"
        new_template = V113 / "template.docx"
        new_manifest_path = V113 / "manifest.yaml"
        self.assertEqual(sha256(old_template), EXPECTED_V112_TEMPLATE_SHA256)
        self.assertEqual(sha256(old_manifest), EXPECTED_V112_MANIFEST_SHA256)

        manifest = yaml.safe_load(new_manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["template"]["version"], "1.1.3")
        self.assertEqual(manifest["template"]["base_manifest"], "v1.1.2/manifest.yaml")
        self.assertEqual(manifest["template"]["base_template"], "v1.1.2/template.docx")
        self.assertEqual(
            manifest["template"]["patches"]["evaluation_table_column_widths_twips"],
            OPTION_B_WIDTHS_TWIPS,
        )
        self.assertEqual(
            manifest["template"]["patches"]["evaluation_table_pagination"],
            {"keep_rows_together": True},
        )
        self.assertEqual(manifest["fingerprint"]["sha256"], sha256(new_template))
        self.assertEqual(manifest["fingerprint"]["value"], sha256(new_template))
        self.assertEqual(sha256(old_template), EXPECTED_V112_TEMPLATE_SHA256)

        _old_document, old_table = _evaluation_table(old_template)
        new_document, new_table = _evaluation_table(new_template)
        self.assertEqual((len(new_table.rows), len(new_table.columns)), (14, 4))
        self.assertEqual((len(old_table.rows), len(old_table.columns)), (14, 4))
        self.assertEqual(
            [int(node.get(qn("w:w"))) for node in new_table._tbl.tblGrid.findall(qn("w:gridCol"))],
            OPTION_B_WIDTHS_TWIPS,
        )
        self.assertEqual(
            int(new_table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))),
            int(old_table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))),
        )
        self.assertEqual(_evaluation_matrix(new_template), _evaluation_matrix(old_template))
        self.assertEqual(_merge_signature(new_template), _merge_signature(old_template))
        self.assertEqual(_bookmark_signature(new_template), _bookmark_signature(old_template))
        self.assertTrue(
            all(row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for row in new_table.rows)
        )

        validation = _run(
            LESSON / "scripts" / "validate_template.py",
            "--template",
            str(new_template),
            "--manifest",
            str(new_manifest_path),
            "--json",
        )
        self.assertEqual(validation.returncode, 0, validation.stderr or validation.stdout)
        report = json.loads(validation.stdout)
        self.assertEqual(report["template_version"], "1.1.3")
        self.assertEqual(report["checks"]["evaluation_table"], {"rows": 14, "columns": 4})
        self.assertEqual(
            report["checks"]["evaluation_table_widths_twips"]["actual_grid"],
            OPTION_B_WIDTHS_TWIPS,
        )
        self.assertTrue(report["checks"]["evaluation_table_widths_twips"]["every_row_matches"])
        self.assertEqual(
            report["checks"]["evaluation_table_pagination"],
            {
                "expected_keep_rows_together": True,
                "actual_rows_kept_together": True,
            },
        )

        with tempfile.TemporaryDirectory(prefix="issue71-v113-rebuild-") as temp_name:
            rebuilt = Path(temp_name) / "template-v1.1.3.docx"
            result = _run(
                LESSON / "scripts" / "build_template_patch.py",
                str(old_template),
                str(rebuilt),
                "--manifest",
                str(new_manifest_path),
            )
            self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
            self.assertEqual(sha256(rebuilt), sha256(new_template))

    def test_same_issue71_input_preserves_the_entire_evaluation_table_across_versions(self) -> None:
        self.assertEqual(sha256(ISSUE71_INPUT), "FD9131E1E74138BD02FB0065D77BA1334DAF8F5CDC615067F25CB105DC29404A")
        source = json.loads(ISSUE71_INPUT.read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory(prefix="issue71-v113-generation-") as temp_name:
            folder = Path(temp_name)
            old_output = folder / "output-v112"
            new_output = folder / "output-v113"
            old_report = _generate(ISSUE71_INPUT, old_output)
            new_report = _generate(ISSUE71_INPUT, new_output, V113 / "template.docx")

            self.assertEqual(old_report["template_version"], "1.1.2")
            self.assertEqual(new_report["template_version"], "1.1.3")
            for report in (old_report, new_report):
                self.assertEqual(report["required_anchor_count"], 70)
                self.assertEqual(report["preserved_anchor_count"], 70)

            old_files = sorted(old_output.glob("*.docx"))
            new_files = sorted(new_output.glob("*.docx"))
            self.assertEqual((len(old_files), len(new_files)), (1, 1))
            self.assertEqual(_evaluation_matrix(old_files[0]), _evaluation_matrix(new_files[0]))
            self.assertEqual(_bookmark_signature(old_files[0]), _bookmark_signature(new_files[0]))
            self.assertEqual(_merge_signature(old_files[0]), _merge_signature(new_files[0]))

            _document, table = _evaluation_table(new_files[0])
            score_cells = [table.cell(row, 2).text.strip() for row in range(1, 14)]
            remarks = [table.cell(row, 3).text for row in range(1, 14)]
            expected_remarks = list(source["lessons"][0]["evaluation"]["remarks"].values())
            self.assertEqual(set(remarks), set(expected_remarks))
            self.assertEqual(sum((float(value) for value in score_cells), 0), source["lessons"][0]["evaluation"]["score"])
            # The committed Issue #71 fixture is score 88.5, so its 25-point
            # rubric cell is 21.5. It does not contain 23.5; keep that fixture
            # untouched and cover the explicitly requested 23.5 in a separate test.
            self.assertIn("21.5", score_cells)
            self.assertNotIn("23.5", score_cells)

    def test_v113_keeps_a_235_score_cell_intact(self) -> None:
        # This transient legacy test case exercises the specified numeric cell
        # without editing the committed Issue #71 score or its provenance.
        fixture = ROOT / "tests" / "fixtures" / "lesson-plan-input.json"
        payload = json.loads(fixture.read_text(encoding="utf-8"))
        payload["lessons"][0]["evaluation"]["score"] = 90.5
        with tempfile.TemporaryDirectory(prefix="lesson-v113-score-235-") as temp_name:
            folder = Path(temp_name)
            source = folder / "score-90.5.json"
            source.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            output = folder / "output"
            report = _generate(source, output, V113 / "template.docx", legacy=True)
            self.assertEqual(report["template_version"], "1.1.3")
            artifacts = json.loads((output / "artifact-manifest.json").read_text(encoding="utf-8"))["artifacts"]
            first = next(item for item in artifacts if item["lesson_id"] == "L01")
            path = output / first["final_docx_path"]
            _document, table = _evaluation_table(path)
            self.assertEqual(table.cell(11, 2).text.strip(), "23.5")


if __name__ == "__main__":
    unittest.main()
