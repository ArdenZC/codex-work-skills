from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
LESSON = ROOT / "教案生成器" / "lesson-plan-docx-generator"
V112 = LESSON / "assets" / "templates" / "lesson-plan" / "v1.1.2"
V114 = LESSON / "assets" / "templates" / "lesson-plan" / "v1.1.4"


class Issue75TemplateV114Tests(unittest.TestCase):
    def test_optional_v114_fits_all_fixed_tables_inside_the_page_body(self) -> None:
        result = subprocess.run(
            [
                sys.executable,
                str(LESSON / "scripts" / "validate_template.py"),
                "--template",
                str(V114 / "template.docx"),
                "--manifest",
                str(V114 / "manifest.yaml"),
                "--json",
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        report = json.loads(result.stdout)
        self.assertEqual(report["template_version"], "1.1.4")
        self.assertEqual(report["checks"]["tables_within_body"]["overwide_tables"], [])
        self.assertGreater(
            report["checks"]["tables_within_body"]["body_width_emu"],
            9918 * 635,
        )
        self.assertTrue(report["checks"]["repeated_column_label_guide"]["declared"])
        self.assertTrue(report["checks"]["reflection_compact_typography"]["natural_row_heights"])
        self.assertTrue(report["checks"]["implementation_table_pagination"]["matches"])
        self.assertTrue(report["checks"]["title_east_asia_font"]["matches"])

    def test_template_patch_builder_applies_readable_implementation_layout(self) -> None:
        with tempfile.TemporaryDirectory(prefix="issue75-v114-patch-") as temp_dir:
            target = Path(temp_dir) / "template.docx"
            result = subprocess.run(
                [
                    sys.executable,
                    str(LESSON / "scripts" / "build_template_patch.py"),
                    str(V114 / "template.docx"),
                    str(target),
                    "--manifest",
                    str(V114 / "manifest.yaml"),
                ],
                cwd=ROOT,
                capture_output=True,
                text=True,
                encoding="utf-8",
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
            document = Document(target)
            title_fonts = document.paragraphs[0].runs[0]._r.rPr.rFonts
            self.assertEqual(title_fonts.get(qn("w:eastAsia")), "仿宋")
            self.assertIsNone(title_fonts.get(qn("w:eastAsiaTheme")))
            table = document.tables[1]
            widths = [int(cell.tcPr.tcW.w) for cell in table.rows[0]._tr.tc_lst]
            self.assertEqual(widths, [1129, 2500, 2100, 2000, 2189])
            content_runs = [
                run
                for row in table.rows[1:]
                for cell in row._tr.tc_lst[1:]
                for run in cell.findall(".//" + qn("w:r"))
                if "".join(node.text or "" for node in run.findall(".//" + qn("w:t"))).strip()
            ]
            self.assertTrue(content_runs)
            self.assertTrue(
                all(
                    run.find(qn("w:rPr")) is not None
                    and run.find(qn("w:rPr")).find(qn("w:sz")) is not None
                    and run.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")) == "20"
                    and run.find(qn("w:rPr")).find(qn("w:szCs")) is not None
                    and run.find(qn("w:rPr")).find(qn("w:szCs")).get(qn("w:val")) == "20"
                    for run in content_runs
                )
            )
            keep_next_rows = []
            for row_index, row in enumerate(table.rows):
                paragraphs = row._tr.findall(".//" + qn("w:p"))
                flags = [
                    paragraph.find(qn("w:pPr")) is not None
                    and paragraph.find(qn("w:pPr")).find(qn("w:keepNext")) is not None
                    for paragraph in paragraphs
                ]
                if flags and all(flags):
                    keep_next_rows.append(row_index)
                else:
                    self.assertFalse(any(flags), f"row {row_index} has inconsistent keep-with-next formatting")
            self.assertEqual(keep_next_rows, [9, 10, 11])

    def test_default_template_remains_v112(self) -> None:
        self.assertTrue((V112 / "template.docx").is_file())
        self.assertTrue((V112 / "manifest.yaml").is_file())
        package_common = (LESSON / "scripts" / "package_common.py").read_text(encoding="utf-8")
        self.assertIn('"v1.1.2" / "manifest.yaml"', package_common)


if __name__ == "__main__":
    unittest.main()
